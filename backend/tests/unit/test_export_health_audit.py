"""
Unit tests for export, health, and audit modules.

Tests cover:
- export/service.py: rows_to_csv and rows_to_xlsx (pure functions, no DB needed)
- health/service.py: get_system_health and get_ingestion_health with mocked DB
- audit/service.py: log_action and get_audit_logs with mocked DB
"""
import csv
import io
from datetime import datetime, timezone
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from fastapi import HTTPException

from app.modules.export.service import rows_to_csv, rows_to_xlsx
from app.modules.health.service import get_ingestion_health, get_system_health
from app.modules.audit.service import get_audit_logs, log_action
from app.modules.audit.models import AuditLog


# ---------------------------------------------------------------------------
# Export service — pure function tests (no DB, no mocks needed)
# ---------------------------------------------------------------------------


class TestRowsToCsv:
    def test_empty_rows_returns_empty_bytes(self):
        result = rows_to_csv([])
        assert result == b""

    def test_single_row(self):
        rows = [{"PATID": "P001", "SEX": "M"}]
        result = rows_to_csv(rows)
        assert isinstance(result, bytes)
        content = result.decode("utf-8")
        reader = csv.DictReader(io.StringIO(content))
        parsed = list(reader)
        assert len(parsed) == 1
        assert parsed[0]["PATID"] == "P001"
        assert parsed[0]["SEX"] == "M"

    def test_multiple_rows(self):
        rows = [
            {"PATID": "P001", "SEX": "M", "RACE": "05"},
            {"PATID": "P002", "SEX": "F", "RACE": "03"},
        ]
        result = rows_to_csv(rows)
        content = result.decode("utf-8")
        reader = csv.DictReader(io.StringIO(content))
        parsed = list(reader)
        assert len(parsed) == 2
        assert parsed[1]["PATID"] == "P002"

    def test_headers_are_included(self):
        rows = [{"COL_A": 1, "COL_B": 2}]
        result = rows_to_csv(rows)
        content = result.decode("utf-8")
        assert "COL_A" in content
        assert "COL_B" in content

    def test_numeric_values_are_exported(self):
        rows = [{"count": 150, "avg": 34.5}]
        result = rows_to_csv(rows)
        content = result.decode("utf-8")
        assert "150" in content
        assert "34.5" in content

    def test_none_values_exported_as_empty(self):
        rows = [{"PATID": "P001", "BIRTH_DATE": None}]
        result = rows_to_csv(rows)
        content = result.decode("utf-8")
        reader = csv.DictReader(io.StringIO(content))
        parsed = list(reader)
        assert parsed[0]["BIRTH_DATE"] == ""

    def test_output_is_utf8_encoded(self):
        rows = [{"name": "José"}]
        result = rows_to_csv(rows)
        assert isinstance(result, bytes)
        decoded = result.decode("utf-8")
        assert "José" in decoded


class TestRowsToXlsx:
    def test_empty_rows_returns_valid_xlsx(self):
        result = rows_to_xlsx([])
        assert isinstance(result, bytes)
        assert len(result) > 0
        # Should be valid xlsx — openpyxl can load it
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.max_row == 1  # only header row or empty

    def test_single_row_with_headers(self):
        rows = [{"PATID": "P001", "SEX": "M"}]
        result = rows_to_xlsx(rows)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        # Row 1 = headers, row 2 = data
        assert ws.cell(row=1, column=1).value in ["PATID", "SEX"]
        assert ws.max_row == 2

    def test_multiple_rows(self):
        rows = [
            {"PATID": "P001", "SEX": "M"},
            {"PATID": "P002", "SEX": "F"},
            {"PATID": "P003", "SEX": "M"},
        ]
        result = rows_to_xlsx(rows)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        # Row 1 = headers, rows 2-4 = data
        assert ws.max_row == 4

    def test_column_order_matches_dict_keys(self):
        rows = [{"FIRST": "A", "SECOND": "B", "THIRD": "C"}]
        result = rows_to_xlsx(rows)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        headers = [ws.cell(row=1, column=i).value for i in range(1, 4)]
        assert headers == ["FIRST", "SECOND", "THIRD"]

    def test_numeric_values_preserved(self):
        rows = [{"count": 42, "rate": 3.14}]
        result = rows_to_xlsx(rows)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        values = [ws.cell(row=2, column=i).value for i in range(1, 3)]
        assert 42 in values
        assert 3.14 in values


# ---------------------------------------------------------------------------
# Health service tests
# ---------------------------------------------------------------------------


class TestGetSystemHealth:
    @pytest.mark.asyncio
    async def test_healthy_when_db_query_succeeds(self):
        db = AsyncMock()
        db.execute = AsyncMock()  # Doesn't raise — db is healthy

        result = await get_system_health(db)

        assert result["status"] == "healthy"
        assert result["components"]["database"] == "healthy"
        assert result["components"]["api"] == "healthy"

    @pytest.mark.asyncio
    async def test_degraded_when_db_query_fails(self):
        db = AsyncMock()
        db.execute.side_effect = Exception("Connection refused")

        result = await get_system_health(db)

        assert result["status"] == "degraded"
        assert "unhealthy" in result["components"]["database"]
        assert "Connection refused" in result["components"]["database"]

    @pytest.mark.asyncio
    async def test_api_always_reported_healthy(self):
        """API component is always healthy since if we get here, the API is running."""
        db = AsyncMock()
        db.execute.side_effect = Exception("DB down")

        result = await get_system_health(db)
        assert result["components"]["api"] == "healthy"


class TestGetIngestionHealth:
    @pytest.mark.asyncio
    async def test_returns_stats_from_db(self):
        db = AsyncMock()
        rows = [("completed", 50), ("failed", 3), ("running", 1)]
        result_mock = MagicMock()
        result_mock.fetchall.return_value = rows
        db.execute.return_value = result_mock

        result = await get_ingestion_health(db)
        stats = result["ingestion_run_stats"]
        assert stats["completed"] == 50
        assert stats["failed"] == 3
        assert stats["running"] == 1

    @pytest.mark.asyncio
    async def test_returns_empty_stats_on_db_error(self):
        db = AsyncMock()
        db.execute.side_effect = Exception("Table missing")

        result = await get_ingestion_health(db)
        assert result["ingestion_run_stats"] == {}


# ---------------------------------------------------------------------------
# Audit service tests
# ---------------------------------------------------------------------------


def _make_audit_log(
    log_id: int = 1,
    action: str = "login",
    user_id: int = 1,
) -> AuditLog:
    entry = AuditLog()
    entry.id = log_id
    entry.action = action
    entry.user_id = user_id
    entry.resource_type = None
    entry.resource_id = None
    entry.details = None
    entry.ip_address = "127.0.0.1"
    entry.timestamp = datetime.now(timezone.utc)
    return entry


class TestLogAction:
    @pytest.mark.asyncio
    async def test_creates_audit_log_entry(self):
        db = AsyncMock()
        db.commit = AsyncMock()
        db.refresh = AsyncMock()

        await log_action(
            db,
            action="query_executed",
            user_id=5,
            resource_type="query",
            resource_id="42",
            details="SELECT * FROM DEMOGRAPHIC",
            ip_address="192.168.1.1",
        )

        db.add.assert_called_once()
        added = db.add.call_args[0][0]
        assert added.action == "query_executed"
        assert added.user_id == 5
        assert added.resource_type == "query"
        assert added.resource_id == "42"
        assert added.ip_address == "192.168.1.1"
        db.commit.assert_called_once()

    @pytest.mark.asyncio
    async def test_log_action_with_only_required_fields(self):
        db = AsyncMock()
        db.commit = AsyncMock()
        db.refresh = AsyncMock()

        await log_action(db, action="login")

        db.add.assert_called_once()
        added = db.add.call_args[0][0]
        assert added.action == "login"
        assert added.user_id is None
        assert added.resource_type is None
        assert added.ip_address is None

    @pytest.mark.asyncio
    async def test_commit_is_always_called(self):
        db = AsyncMock()
        db.commit = AsyncMock()
        db.refresh = AsyncMock()

        await log_action(db, action="export_csv", user_id=1)
        db.commit.assert_called_once()


class TestGetAuditLogs:
    @pytest.mark.asyncio
    async def test_returns_list_of_logs(self):
        log1 = _make_audit_log(log_id=1, action="login")
        log2 = _make_audit_log(log_id=2, action="query_executed")
        result_mock = MagicMock()
        result_mock.scalars.return_value.all.return_value = [log2, log1]

        db = AsyncMock()
        db.execute.return_value = result_mock

        logs = await get_audit_logs(db)
        assert len(logs) == 2
        assert logs[0].action == "query_executed"

    @pytest.mark.asyncio
    async def test_returns_empty_list_when_no_logs(self):
        result_mock = MagicMock()
        result_mock.scalars.return_value.all.return_value = []

        db = AsyncMock()
        db.execute.return_value = result_mock

        logs = await get_audit_logs(db)
        assert logs == []

    @pytest.mark.asyncio
    async def test_filtering_by_user_id_passes_query_to_db(self):
        """Verify that a DB execute is called (filtering logic is in the query, tested here indirectly)."""
        result_mock = MagicMock()
        result_mock.scalars.return_value.all.return_value = []

        db = AsyncMock()
        db.execute.return_value = result_mock

        await get_audit_logs(db, user_id=42)
        db.execute.assert_called_once()

    @pytest.mark.asyncio
    async def test_filtering_by_action_passes_query_to_db(self):
        result_mock = MagicMock()
        result_mock.scalars.return_value.all.return_value = []

        db = AsyncMock()
        db.execute.return_value = result_mock

        await get_audit_logs(db, action="login")
        db.execute.assert_called_once()
